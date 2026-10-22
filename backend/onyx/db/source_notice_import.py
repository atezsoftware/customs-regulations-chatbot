"""Manifest-only planning and append-only application of repealed source notices."""

from __future__ import annotations

import hashlib
import json
import re
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Literal, cast
from uuid import uuid4

from pydantic import BaseModel, Field, JsonValue
from sqlalchemy import text
from sqlalchemy.orm import Session

from onyx.regulatory.source_notice_identity import (
    NoticeIdentity,
    normalize_notice_alias,
    notice_identity_queries,
    notice_title_alias,
)

MANIFEST_VERSION = 1


def guard_notice_import_statement(statement: str) -> None:
    """Reject all SQL writes except a single INSERT into the new notice table."""
    normalized = statement.strip().rstrip(";").strip()
    if ";" in normalized or "--" in normalized or "/*" in normalized:
        raise ValueError(
            "Multi-statement/comment SQL is not allowed by the notice importer"
        )
    if re.search(
        r"\b(?:UPDATE|DELETE|TRUNCATE|ALTER|DROP|CREATE|MERGE|CALL|DO|COPY|GRANT|REVOKE)\b",
        normalized,
        re.IGNORECASE,
    ):
        raise ValueError("Data-changing SQL tokens are forbidden")
    if re.search(r"\b(?:nextval|setval)\s*\(", normalized, re.IGNORECASE):
        raise ValueError("Sequence-changing functions are forbidden")
    for function in re.findall(r"\b(pg_[a-z_]+)\s*\(", normalized, re.IGNORECASE):
        if function.lower() != "pg_advisory_xact_lock":
            raise ValueError("Unapproved PostgreSQL function in notice import")
    if re.match(r"^(?:SELECT|SHOW)\b", normalized, re.IGNORECASE):
        if re.search(
            r"\b(?:INTO|UPDATE|DELETE|INSERT|ALTER|CREATE|DROP|TRUNCATE)\b",
            normalized,
            re.IGNORECASE,
        ):
            raise ValueError("Writing SELECT statements are forbidden")
        return
    if re.match(
        r"^INSERT\s+INTO\s+regulatory_source_notice\s*\(", normalized, re.IGNORECASE
    ):
        if re.search(r"\bON\s+CONFLICT\b", normalized, re.IGNORECASE):
            raise ValueError("Notice imports never update or replace existing rows")
        return
    raise ValueError(
        "Only SELECT/SHOW and INSERT INTO regulatory_source_notice are allowed"
    )


class ExcelNoticeRow(BaseModel):
    excel_row: int
    raw_values: list[JsonValue]
    category: str
    catalog_title: str
    title_column: str


class ExistingNoticeSource(BaseModel):
    id: str
    name: str
    file_id: str | None = None
    status: str | None = None
    legistation_metadata: dict[str, JsonValue] | None = None


class SourceIdentityHead(BaseModel):
    source_id: str
    chunk_id: str
    position: int
    projection_ordinal: int
    heading_path: list[str]
    document_type: str | None = None
    document_number: str | None = None
    document_date: str | None = None
    text_head: str


class NoticeImportTarget(BaseModel):
    id: str
    document_set_id: int
    identity_key: str
    instrument_type: str
    instrument_number: str | None = None
    instrument_year: int | None = None
    instrument_date: date | None = None
    catalog_title: str
    aliases: list[str]
    legistation_metadata: dict[str, JsonValue]
    provenance: dict[str, JsonValue]
    content_state: str = "metadata_only"
    user_file_id: str | None = None
    revision: int = 1


class NoticeRowDecision(BaseModel):
    excel_rows: list[int]
    identity_key: str | None = None
    classification: str
    reason: str
    catalog_title: str
    existing_source_ids: list[str] = Field(default_factory=list)
    raw_records: list[ExcelNoticeRow]


class SourceNoticeManifest(BaseModel):
    schema_version: int = MANIFEST_VERSION
    workbook_sha256: str
    source_path: str
    sheet: str
    document_set_id: int
    inventory_fingerprint: str
    inventory_count: int
    identity_heads_fingerprint: str | None = None
    identity_heads_count: int = 0
    raw_row_count: int
    targets: list[NoticeImportTarget]
    decisions: list[NoticeRowDecision]
    admission_policy: Literal["structured_only", "catalogue_observations"] = (
        "structured_only"
    )
    prior_manifest_sha256: str | None = None
    registry_fingerprint: str | None = None
    baseline_notices: list[NoticeImportTarget] = Field(default_factory=list)


def canonical_digest(value: object) -> str:
    data = json.dumps(
        value, sort_keys=True, ensure_ascii=False, separators=(",", ":"), default=str
    )
    return hashlib.sha256(data.encode()).hexdigest()


def notice_manifest_digests(manifest: SourceNoticeManifest) -> set[str]:
    """Retain verification of previously frozen manifests without new optional fields."""
    payload = manifest.model_dump(mode="json")
    digests = {canonical_digest(payload)}
    if (
        manifest.admission_policy == "structured_only"
        and not manifest.baseline_notices
        and manifest.registry_fingerprint is None
        and manifest.prior_manifest_sha256 is None
    ):
        for key in (
            "admission_policy",
            "prior_manifest_sha256",
            "registry_fingerprint",
            "baseline_notices",
        ):
            payload.pop(key)
        digests.add(canonical_digest(payload))
    return digests


def notice_registry_fingerprint(notices: list[NoticeImportTarget]) -> str:
    return canonical_digest(
        [
            notice.model_dump(mode="json")
            for notice in sorted(notices, key=lambda notice: notice.identity_key)
        ]
    )


def inventory_fingerprint(sources: list[ExistingNoticeSource]) -> str:
    return canonical_digest(
        [
            source.model_dump(mode="json")
            for source in sorted(sources, key=lambda source: source.id)
        ]
    )


def read_notice_workbook(
    path: Path, *, expected_sha256: str
) -> tuple[str, list[ExcelNoticeRow]]:
    from openpyxl import load_workbook

    if hashlib.sha256(path.read_bytes()).hexdigest() != expected_sha256:
        raise ValueError("Workbook SHA256 differs from the reviewed source")
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if len(workbook.worksheets) != 1:
            raise ValueError("Expected one reviewed worksheet")
        sheet = workbook.worksheets[0]
        if sheet.max_column != 4:
            raise ValueError("Expected the reviewed four-column layout")
        rows: list[ExcelNoticeRow] = []
        for row_number, cells in enumerate(sheet.iter_rows(min_col=1, max_col=4), 1):
            if any(cell.data_type == "f" for cell in cells):
                raise ValueError(f"Formula in input row {row_number}")
            raw: list[JsonValue] = []
            for cell in cells:
                value = cell.value
                if isinstance(value, (datetime, date)):
                    raw.append(value.isoformat())
                elif isinstance(value, (str, int, float, bool)) or value is None:
                    raw.append(value)
                else:
                    raise ValueError(f"Unsupported input cell in row {row_number}")
            if not any(value is not None for value in raw):
                continue
            category = raw[0]
            if not isinstance(category, str):
                raise ValueError(f"Missing category in row {row_number}")
            title_column = (
                "D"
                if category
                in (
                    "Genelgeler",
                    "Kararname ve Mahkeme Kararları",
                    "Tasarruflu Yazılar",
                )
                else "C"
            )
            title = raw[3 if title_column == "D" else 2]
            if not isinstance(title, str) or not title.strip():
                raise ValueError(f"Missing catalog title in row {row_number}")
            rows.append(
                ExcelNoticeRow(
                    excel_row=row_number,
                    raw_values=raw,
                    category=category,
                    catalog_title=title,
                    title_column=title_column,
                )
            )
        return sheet.title, rows
    finally:
        workbook.close()


def _positive_number(value: JsonValue) -> str | None:
    if isinstance(value, bool):
        return None
    raw = str(value)
    if not re.fullmatch(r"[0-9]{1,15}", raw) or int(raw) <= 0:
        return None
    return str(int(raw))


def row_identity(row: ExcelNoticeRow) -> tuple[NoticeIdentity | None, str]:
    first, second = row.raw_values[1:3]
    if row.category == "Diğer Kanunlar":
        number = _positive_number(first)
        return (
            (NoticeIdentity("law", number), "structured_law_number")
            if number
            else (None, "invalid_law_number")
        )
    if row.category in ("Genelgeler", "Kararname ve Mahkeme Kararları"):
        year, number = _positive_number(first), _positive_number(second)
        if not year or not number or not 1800 <= int(year) <= 2099:
            return None, "missing_or_unstructured_year_number"
        kind = "circular" if row.category == "Genelgeler" else "decision"
        if kind == "circular" and int(number) > 999:
            return None, "unverified_circular_number_or_external_registry_number"
        identity = NoticeIdentity(kind, number, int(year))
        # A conflicting explicit own year/number in the description needs review.
        pairs = {
            (int(match[0]), str(int(match[1])))
            for match in re.findall(
                r"(?<!\d)((?:18|19|20)\d{2})\s*[/–—-]\s*(\d+)(?!\d)", row.catalog_title
            )
        }
        if pairs and (int(year), number) not in pairs:
            return None, "structured_identity_conflicts_with_catalog_text"
        return identity, "structured_type_year_number"
    if row.category == "Tasarruflu Yazılar":
        number = _positive_number(second)
        if not isinstance(first, str) or not number:
            return None, "missing_or_unstructured_letter_date_number"
        try:
            issued = datetime.fromisoformat(first).date()
        except ValueError:
            return None, "invalid_letter_date"
        return NoticeIdentity(
            "administrative_letter", number, instrument_date=issued
        ), "structured_letter_date_number"
    return None, "unverified_title_or_series_identity"


_STOPWORDS = frozenset(
    "mulga hakkinda dair iliskin ile ve bir no sayili seri teblig tebligi genelge karar karari kanun kanunu yonetmelik yonetmeligi yazilar kanunlar diger mevzuati".split()
)


def _title_tokens(title: str) -> frozenset[str]:
    return frozenset(
        token
        for token in normalize_notice_alias(title).split()
        if len(token) > 2 and not token.isdigit() and token not in _STOPWORDS
    )


@dataclass
class _CatalogIndex:
    exact_keys: dict[str, set[str]]
    number_sources: dict[str, set[str]]
    pair_sources: dict[tuple[str, str], set[str]]
    token_sources: dict[str, set[str]]
    tokens: dict[str, frozenset[str]]
    exact_title_sources: dict[frozenset[str], set[str]]

    @classmethod
    def build(cls, sources: list[ExistingNoticeSource]) -> _CatalogIndex:
        index = cls(
            defaultdict(set),
            defaultdict(set),
            defaultdict(set),
            defaultdict(set),
            {},
            defaultdict(set),
        )
        for source in sources:
            metadata = source.legistation_metadata or {}
            fields = [
                source.name,
                *(
                    value
                    for key, value in metadata.items()
                    if key in ("official_title", "subject") and isinstance(value, str)
                ),
            ]
            for field in fields:
                field_tokens = _title_tokens(field)
                if field_tokens:
                    index.exact_title_sources[field_tokens].add(source.id)
            index.tokens[source.id] = frozenset().union(
                *(_title_tokens(value) for value in fields)
            )
            for token in index.tokens[source.id]:
                index.token_sources[token].add(source.id)
            for value in fields:
                for key in notice_identity_queries(value):
                    index.exact_keys[key].add(source.id)
                for number in re.findall(r"\d+", value):
                    index.number_sources[str(int(number))].add(source.id)
                # Candidate-only matching: unrelated year/number pairs cause a hold,
                # never an asserted existing identity or a write to that source.
                for year, number in re.findall(
                    r"(?<!\d)((?:18|19|20)\d{2})\s*[/_–—-]\s*(\d+)(?!\d)", value
                ):
                    index.pair_sources[(year, str(int(number)))].add(source.id)
        return index

    def candidates(
        self, row: ExcelNoticeRow, identity: NoticeIdentity
    ) -> tuple[set[str], set[str]]:
        exact = self.exact_keys.get(identity.key, set())
        if identity.instrument_type == "circular":
            possible = set(
                self.pair_sources.get((str(identity.year), identity.number), set())
            )
        else:
            possible = set(self.number_sources.get(identity.number, set()))
        tokens = _title_tokens(row.catalog_title)
        possible.update(self.exact_title_sources.get(tokens, set()))
        candidates: set[str] = set()
        for token in tokens:
            candidates.update(self.token_sources.get(token, set()))
        for source_id in candidates:
            overlap = tokens & self.tokens[source_id]
            if (
                len(tokens) >= 3
                and len(overlap) >= 3
                and len(overlap) / len(tokens) >= 0.7
            ):
                possible.add(source_id)
        return set(exact), possible - set(exact)


def build_source_notice_manifest(
    *,
    workbook_path: Path,
    workbook_sha256: str,
    sources: list[ExistingNoticeSource],
    document_set_id: int,
    identity_heads: list[SourceIdentityHead] | None = None,
) -> SourceNoticeManifest:
    sheet, rows = read_notice_workbook(workbook_path, expected_sha256=workbook_sha256)
    raw_groups: dict[str, list[ExcelNoticeRow]] = defaultdict(list)
    for row in rows:
        raw_groups[canonical_digest(row.raw_values)].append(row)
    groups_by_key: dict[str, list[list[ExcelNoticeRow]]] = defaultdict(list)
    missing: list[list[ExcelNoticeRow]] = []
    for group in raw_groups.values():
        identity, _reason = row_identity(group[0])
        if identity is None:
            missing.append(group)
        else:
            groups_by_key[identity.key].append(group)
    index = _CatalogIndex.build(sources)
    head_keys: dict[str, set[str]] = defaultdict(set)
    for head in identity_heads or []:
        for key in identity_head_veto_keys(head):
            head_keys[key].add(head.source_id)
    decisions: list[NoticeRowDecision] = []
    targets: list[NoticeImportTarget] = []
    for group in missing:
        _identity, reason = row_identity(group[0])
        decisions.append(
            NoticeRowDecision(
                excel_rows=[row.excel_row for row in group],
                classification="ambiguous",
                reason=reason,
                catalog_title=group[0].catalog_title,
                raw_records=group,
            )
        )
    for key, groups in sorted(groups_by_key.items()):
        group = [row for raw_group in groups for row in raw_group]
        row = group[0]
        identity, _reason = row_identity(row)
        assert identity is not None
        exact, possible = index.candidates(row, identity)
        classification, reason = (
            "new",
            "structured_identity_absent_from_current_catalog",
        )
        head_candidates = head_keys.get(identity.key, set())
        candidates = exact | possible | head_candidates
        if len(groups) > 1:
            classification, reason = (
                "ambiguous",
                "multiple_catalog_records_share_unqualified_identity",
            )
        elif exact:
            classification, reason = (
                "existing",
                "typed_identity_in_existing_source_catalog",
            )
        elif head_candidates:
            classification, reason = (
                "ambiguous",
                "canonical_identity_head_candidate_requires_review",
            )
        elif possible:
            classification, reason = (
                "ambiguous",
                "existing_number_or_title_candidate_requires_review",
            )
        decision = NoticeRowDecision(
            excel_rows=[item.excel_row for item in group],
            identity_key=key,
            classification=classification,
            reason=reason,
            catalog_title=row.catalog_title,
            existing_source_ids=sorted(candidates),
            raw_records=group,
        )
        decisions.append(decision)
        if classification != "new":
            continue
        targets.append(
            NoticeImportTarget(
                id=str(uuid4()),
                document_set_id=document_set_id,
                identity_key=key,
                instrument_type=identity.instrument_type,
                instrument_number=identity.number,
                instrument_year=identity.year,
                instrument_date=identity.instrument_date,
                catalog_title=row.catalog_title,
                aliases=list(identity.aliases),
                legistation_metadata={
                    "validity_status": "Mülga",
                    "repeal_note": row.catalog_title,
                    "validity_source": workbook_path.name,
                },
                provenance={
                    "source_filename": workbook_path.name,
                    "source_sha256": workbook_sha256,
                    "sheet": sheet,
                    "excel_rows": [item.excel_row for item in group],
                    "raw_records": [item.model_dump(mode="json") for item in group],
                    "identity_basis": _reason,
                    "annotation_only": True,
                },
            )
        )
    return SourceNoticeManifest(
        workbook_sha256=workbook_sha256,
        source_path=str(workbook_path),
        sheet=sheet,
        document_set_id=document_set_id,
        inventory_fingerprint=inventory_fingerprint(sources),
        inventory_count=len(sources),
        identity_heads_fingerprint=canonical_digest(
            [head.model_dump(mode="json") for head in identity_heads]
        )
        if identity_heads is not None
        else None,
        identity_heads_count=len(identity_heads or []),
        raw_row_count=len(rows),
        targets=targets,
        decisions=sorted(decisions, key=lambda item: min(item.excel_rows)),
    )


def identity_head_veto_keys(head: SourceIdentityHead) -> set[str]:
    """Potential own-header matches are vetoes, never proof of source identity."""
    keys: set[str] = set()
    for heading in head.heading_path[:1]:
        # Long sentence/bullet paths are extracted body fragments, not titles.
        if len(heading) <= 180 and not heading.lstrip().startswith(("-", "|", "*")):
            keys.update(notice_identity_queries(heading))
    for line in head.text_head.splitlines()[:12]:
        stripped = line.strip(" #*\t")
        if len(stripped) > 180:
            continue
        if re.match(
            r"^(?:genelge|karar(?:ı|i)?|bkk|kanun|tebliğ|teblig)\b",
            stripped,
            re.IGNORECASE,
        ):
            keys.update(notice_identity_queries(stripped))
        law_number = re.match(
            r"^kanun\s+(?:numarası|numarasi|no)\s*[:.]?\s*(\d{1,7})\b",
            stripped,
            re.IGNORECASE,
        )
        if law_number:
            keys.add(NoticeIdentity("law", law_number[1]).key)
    candidate_numbers: set[str] = set()
    candidate_dates: set[date] = set()
    raw_date = head.document_date or ""
    try:
        candidate_dates.add(date.fromisoformat(raw_date[:10]))
    except ValueError:
        pass
    raw_number = (head.document_number or "").strip()
    number_match = re.fullmatch(r"(?:(?:18|19|20)\d{2}[/–—-])?(\d{1,15})", raw_number)
    if number_match:
        candidate_numbers.add(str(int(number_match[1])))
    for line in head.text_head.splitlines()[:12]:
        stripped = normalize_notice_alias(line)
        if not re.match(r"^(?:sayi|sayisi|tarih|tarihi)\b", stripped):
            continue
        if re.match(r"^(?:sayi|sayisi)\b", stripped):
            header_number = re.match(
                r"^\s*(?:Sayı|Sayisi|Sayi|Sayısı)\s*[:.]?\s*(?:(?:18|19|20)\d{2}[/–—-])?(\d{1,15})(?:\s|$)",
                line.strip(" #*\t"),
                re.IGNORECASE,
            )
            if header_number:
                candidate_numbers.add(str(int(header_number[1])))
        for day, month, year in re.findall(
            r"(?<!\d)(\d{1,2})[./-](\d{1,2})[./-]((?:19|20)\d{2})(?!\d)", line
        ):
            try:
                candidate_dates.add(date(int(year), int(month), int(day)))
            except ValueError:
                continue
    for number in candidate_numbers:
        for issued in candidate_dates:
            keys.add(
                NoticeIdentity(
                    "administrative_letter", number, instrument_date=issued
                ).key
            )
    # Chunker metadata can come from a cited passage; treat it only as a veto.
    kind = normalize_notice_alias(head.document_type or "")
    number = (head.document_number or "").strip()
    if kind in ("law", "kanun") and re.fullmatch(r"\d{1,7}", number):
        keys.add(NoticeIdentity("law", number).key)
    kinds = {
        "circular": "circular",
        "genelge": "circular",
        "decision": "decision",
        "karar": "decision",
    }
    if kind in kinds:
        pair = re.fullmatch(r"((?:18|19|20)\d{2})[/–—-](\d+)", number)
        if pair:
            keys.add(NoticeIdentity(kinds[kind], pair[2], int(pair[1])).key)
    return keys


def read_source_identity_heads(
    db_session: Session, document_set_id: int
) -> list[SourceIdentityHead]:
    records = db_session.execute(
        text("""
        SELECT u.id::text AS source_id,c.id AS chunk_id,c.position,c.projection_ordinal,
               c.heading_path,c.chunk_metadata->>'document_type' AS document_type,
               c.chunk_metadata->>'document_number' AS document_number,
               c.chunk_metadata->>'document_date' AS document_date,
               left(c.text,1800) AS text_head
        FROM document_set__user_file m
        JOIN user_file u ON u.id=m.user_file_id
        CROSS JOIN LATERAL (
            SELECT id,position,projection_ordinal,heading_path,chunk_metadata,text
            FROM regulatory_chunk r WHERE r.user_file_id=u.id
            ORDER BY r.position,r.projection_ordinal,r.id LIMIT 2
        ) c
        WHERE m.document_set_id=:scope
        ORDER BY u.id,c.position,c.projection_ordinal,c.id
    """),
        {"scope": document_set_id},
    ).mappings()
    return [SourceIdentityHead.model_validate(dict(row)) for row in records]


def read_current_inventory(
    db_session: Session, document_set_id: int
) -> list[ExistingNoticeSource]:
    records = db_session.execute(
        text("""
        SELECT u.id::text AS id,u.name,u.file_id,u.status::text AS status,
               f.file_metadata->'legistation_metadata' AS legistation_metadata
        FROM document_set__user_file m
        JOIN user_file u ON u.id=m.user_file_id
        LEFT JOIN file_record f ON f.file_id=u.file_id
        WHERE m.document_set_id=:document_set_id ORDER BY u.id
    """),
        {"document_set_id": document_set_id},
    ).mappings()
    return [ExistingNoticeSource.model_validate(dict(row)) for row in records]


def apply_source_notice_manifest(
    db_session: Session,
    manifest: SourceNoticeManifest,
    *,
    expected_manifest_sha256: str,
    expected_workbook_sha256: str,
) -> list[str]:
    """INSERT only new notice rows in the caller's transaction; never commit here."""
    if manifest.identity_heads_fingerprint is None:
        raise ValueError("Apply requires the reviewed canonical identity-head audit")
    if manifest.schema_version != MANIFEST_VERSION:
        raise ValueError("Unsupported manifest schema")
    if expected_manifest_sha256 not in notice_manifest_digests(manifest):
        raise ValueError("Manifest digest differs from reviewed plan")
    if (
        manifest.workbook_sha256 != expected_workbook_sha256
        or hashlib.sha256(Path(manifest.source_path).read_bytes()).hexdigest()
        != expected_workbook_sha256
    ):
        raise ValueError("Workbook changed after planning")
    if (
        inventory_fingerprint(
            read_current_inventory(db_session, manifest.document_set_id)
        )
        != manifest.inventory_fingerprint
    ):
        raise ValueError(
            "Existing source catalog changed after planning; repeat dry run"
        )
    if manifest.identity_heads_fingerprint is not None:
        current_heads = read_source_identity_heads(db_session, manifest.document_set_id)
        if (
            canonical_digest([head.model_dump(mode="json") for head in current_heads])
            != manifest.identity_heads_fingerprint
        ):
            raise ValueError(
                "Canonical identity heads changed after planning; repeat dry run"
            )
    # A transaction-scoped importer lock prevents concurrent manifest application.
    db_session.execute(
        text("SELECT pg_advisory_xact_lock(1762751,:scope)"),
        {"scope": manifest.document_set_id},
    )
    if manifest.admission_policy == "catalogue_observations":
        current_registry = read_notice_registry(db_session, manifest.document_set_id)
        expected_baseline = notice_registry_fingerprint(manifest.baseline_notices)
        if manifest.registry_fingerprint != expected_baseline:
            raise ValueError("Manifest registry baseline fingerprint is invalid")
        expected_after = [*manifest.baseline_notices, *manifest.targets]
        if len({item.identity_key for item in expected_after}) != len(expected_after):
            raise ValueError("Manifest registry baseline and new targets overlap")
        if notice_registry_fingerprint(current_registry) not in {
            expected_baseline,
            notice_registry_fingerprint(expected_after),
        }:
            raise ValueError("Notice registry changed after planning; repeat dry run")
    existing_records = (
        db_session.execute(
            text(
                "SELECT id::text AS id,identity_key,catalog_title,provenance FROM regulatory_source_notice WHERE document_set_id=:scope AND identity_key=ANY(:keys)"
            ),
            {
                "scope": manifest.document_set_id,
                "keys": [target.identity_key for target in manifest.targets],
            },
        )
        .mappings()
        .all()
    )
    existing_by_key = {
        str(record["identity_key"]): record for record in existing_records
    }
    pending: list[NoticeImportTarget] = []
    for target in manifest.targets:
        if (
            target.document_set_id != manifest.document_set_id
            or target.content_state != "metadata_only"
            or target.user_file_id is not None
            or target.revision != 1
        ):
            raise ValueError("Manifest target is not a new metadata-only notice")
        existing = existing_by_key.get(target.identity_key)
        if existing is not None:
            prior = cast(dict[str, JsonValue], existing["provenance"])
            if (
                existing["catalog_title"] != target.catalog_title
                or prior.get("source_sha256") != manifest.workbook_sha256
                or prior.get("excel_rows") != target.provenance["excel_rows"]
            ):
                raise ValueError(
                    f"Notice identity already exists with different provenance: {target.identity_key}"
                )
            continue
        pending.append(target)
    for offset in range(0, len(pending), 200):
        _insert_notice_batch(db_session, pending[offset : offset + 200])
    return [target.id for target in pending]


def _insert_notice_batch(
    db_session: Session, targets: list[NoticeImportTarget]
) -> None:
    from uuid import UUID

    from sqlalchemy import insert

    from onyx.db.models import RegulatorySourceNotice

    values = []
    for target in targets:
        record = target.model_dump(mode="python")
        record["id"] = UUID(target.id)
        values.append(record)
    statement = insert(RegulatorySourceNotice).values(values)
    guard_notice_import_statement(str(statement))
    db_session.execute(statement)


def validate_imported_notice_targets(
    manifest: SourceNoticeManifest, actual: list[NoticeImportTarget]
) -> dict[str, JsonValue]:
    """Compare every imported business field, including raw text and provenance."""
    expected_records = [*manifest.baseline_notices, *manifest.targets]
    expected_by_key = {target.identity_key: target for target in expected_records}
    actual_by_key = {target.identity_key: target for target in actual}
    if len(expected_by_key) != len(expected_records) or len(actual_by_key) != len(
        actual
    ):
        raise ValueError("Duplicate notice identity in manifest or stored registry")
    if set(actual_by_key) != set(expected_by_key):
        raise ValueError("Stored notice identity set differs from reviewed manifest")
    observed_rows: set[int] = set()
    for key, expected in expected_by_key.items():
        stored = actual_by_key[key]
        if stored.model_dump(mode="json") != expected.model_dump(mode="json"):
            raise ValueError(
                f"Stored notice fields differ from reviewed manifest: {key}"
            )
        if (
            stored.legistation_metadata.get("validity_status") != "Mülga"
            or stored.content_state != "metadata_only"
            or stored.user_file_id is not None
            or stored.revision != 1
        ):
            raise ValueError("Stored notice is not an unmaterialized repealed record")
        rows = stored.provenance.get("excel_rows")
        if not isinstance(rows, list) or any(
            isinstance(row, bool) or not isinstance(row, int) for row in rows
        ):
            raise ValueError("Stored provenance has invalid source row numbers")
        observed_rows.update(cast(list[int], rows))
    skipped_rows = {
        row
        for decision in manifest.decisions
        if decision.classification not in ("new", "already_registered")
        for row in decision.excel_rows
    }
    skipped_keys = {
        decision.identity_key
        for decision in manifest.decisions
        if decision.classification not in ("new", "already_registered")
        and decision.identity_key is not None
    }
    if observed_rows & skipped_rows or set(actual_by_key) & skipped_keys:
        raise ValueError("A held or existing source row was inserted")
    ordered = [
        actual_by_key[key].model_dump(mode="json") for key in sorted(actual_by_key)
    ]
    return {
        "verified_notice_count": len(actual),
        "new_target_count": len(manifest.targets),
        "unchanged_baseline_count": len(manifest.baseline_notices),
        "verified_business_field_sha256": canonical_digest(ordered),
        "all_repealed": True,
        "all_metadata_only": True,
        "all_user_file_ids_null": True,
        "existing_record_groups_skipped": sum(
            decision.classification == "existing" for decision in manifest.decisions
        ),
        "held_raw_rows_skipped": len(skipped_rows)
        - sum(
            len(decision.excel_rows)
            for decision in manifest.decisions
            if decision.classification == "existing"
        ),
        "skipped_source_row_intersections": 0,
        "skipped_identity_intersections": 0,
    }


def read_notice_registry(
    db_session: Session, document_set_id: int
) -> list[NoticeImportTarget]:
    records = db_session.execute(
        text("""
        SELECT id::text AS id,document_set_id,identity_key,instrument_type,
               instrument_number,instrument_year,instrument_date,catalog_title,
               aliases,legistation_metadata,provenance,content_state,
               user_file_id::text AS user_file_id,revision
        FROM regulatory_source_notice WHERE document_set_id=:scope
        ORDER BY identity_key
    """),
        {"scope": document_set_id},
    ).mappings()
    return [NoticeImportTarget.model_validate(dict(record)) for record in records]


def verify_stored_source_notice_manifest(
    db_session: Session,
    manifest: SourceNoticeManifest,
    *,
    expected_manifest_sha256: str,
) -> dict[str, JsonValue]:
    if expected_manifest_sha256 not in notice_manifest_digests(manifest):
        raise ValueError("Manifest digest differs from reviewed plan")
    actual = read_notice_registry(db_session, manifest.document_set_id)
    result = validate_imported_notice_targets(manifest, actual)
    counts = db_session.execute(
        text(
            "SELECT document_set_id,count(*)::int AS count FROM regulatory_source_notice GROUP BY document_set_id ORDER BY document_set_id"
        )
    ).mappings()
    result["registry_scope_counts"] = [
        {
            "document_set_id": int(record["document_set_id"]),
            "count": int(record["count"]),
        }
        for record in counts
    ]
    return result


def _catalogue_observation_fields(
    row: ExcelNoticeRow,
) -> tuple[str, str | None, int | None, date | None]:
    identity, _basis = row_identity(row)
    if identity is not None:
        return (
            identity.instrument_type,
            identity.number,
            identity.year,
            identity.instrument_date,
        )
    kinds = {
        "Diğer Kanunlar": "law",
        "Genelgeler": "circular",
        "Kararname ve Mahkeme Kararları": "decision",
        "Tasarruflu Yazılar": "administrative_letter",
        "Diğer Yönetmelikler": "regulation",
        "Tebliğler": "communique",
    }
    kind = kinds.get(row.category, "catalogue_observation")
    first, second = row.raw_values[1:3]
    number: str | None = None
    year: int | None = None
    issued: date | None = None
    if row.category == "Diğer Kanunlar":
        number = _positive_number(first)
    elif row.category in ("Genelgeler", "Kararname ve Mahkeme Kararları"):
        raw_year = _positive_number(first)
        year = (
            int(raw_year)
            if raw_year is not None and 1800 <= int(raw_year) <= 2099
            else None
        )
        number = _positive_number(second)
        if (
            number is None
            and isinstance(second, str)
            and re.fullmatch(r"[A-Za-z]+-\d+", second)
        ):
            number = second
    elif row.category == "Tasarruflu Yazılar":
        number = _positive_number(second)
        if isinstance(first, str):
            try:
                issued = datetime.fromisoformat(first).date()
            except ValueError:
                pass
    return kind, number, year, issued


def build_catalogue_observation_manifest(
    *,
    workbook_path: Path,
    workbook_sha256: str,
    sources: list[ExistingNoticeSource],
    identity_heads: list[SourceIdentityHead],
    document_set_id: int,
    prior_manifest: SourceNoticeManifest,
    prior_manifest_sha256: str,
    registry: list[NoticeImportTarget],
) -> SourceNoticeManifest:
    """Admit independent catalogue observations without mutating matched source files."""
    if prior_manifest_sha256 not in notice_manifest_digests(prior_manifest):
        raise ValueError("Prior reviewed manifest digest differs")
    if (
        prior_manifest.workbook_sha256 != workbook_sha256
        or prior_manifest.document_set_id != document_set_id
    ):
        raise ValueError("Prior reviewed manifest belongs to a different source/scope")
    sheet, rows = read_notice_workbook(workbook_path, expected_sha256=workbook_sha256)
    raw_groups: dict[str, list[ExcelNoticeRow]] = defaultdict(list)
    for row in rows:
        raw_groups[canonical_digest(row.raw_values)].append(row)
    prior_records: dict[str, NoticeRowDecision] = {}
    for decision in prior_manifest.decisions:
        for record in decision.raw_records:
            digest = canonical_digest(record.raw_values)
            if digest in prior_records and prior_records[digest] != decision:
                raise ValueError("Prior decisions disagree for the same raw record")
            prior_records[digest] = decision
    if set(prior_records) != set(raw_groups):
        raise ValueError(
            "Current raw Excel records differ from prior reviewed manifest"
        )
    registry_by_key = {notice.identity_key: notice for notice in registry}
    if len(registry_by_key) != len(registry):
        raise ValueError("Duplicate identity in existing notice registry")
    for target in prior_manifest.targets:
        existing = registry_by_key.get(target.identity_key)
        if existing is None or existing.model_dump(mode="json") != target.model_dump(
            mode="json"
        ):
            raise ValueError("Previously imported notice changed or disappeared")
    prior_by_raw = {
        canonical_digest(record.raw_values): target
        for target in prior_manifest.targets
        for record in [
            ExcelNoticeRow.model_validate(item)
            for item in cast(list[JsonValue], target.provenance["raw_records"])
            if isinstance(item, dict)
        ]
    }
    decisions: list[NoticeRowDecision] = []
    targets: list[NoticeImportTarget] = []
    for raw_digest, group in raw_groups.items():
        row = group[0]
        prior_decision = prior_records[raw_digest]
        if prior_decision.classification == "existing":
            decisions.append(
                NoticeRowDecision(
                    excel_rows=[item.excel_row for item in group],
                    identity_key=prior_decision.identity_key,
                    classification="existing",
                    reason="prior_reviewed_confirmed_existing_source_unchanged",
                    catalog_title=row.catalog_title,
                    existing_source_ids=prior_decision.existing_source_ids,
                    raw_records=group,
                )
            )
            continue
        prior_target = prior_by_raw.get(raw_digest)
        if prior_target is not None:
            decisions.append(
                NoticeRowDecision(
                    excel_rows=[item.excel_row for item in group],
                    identity_key=prior_target.identity_key,
                    classification="already_registered",
                    reason="previously_imported_typed_notice_unchanged",
                    catalog_title=row.catalog_title,
                    existing_source_ids=[],
                    raw_records=group,
                )
            )
            continue
        key = "catalog:" + raw_digest
        kind, number, year, issued = _catalogue_observation_fields(row)
        identity, basis = row_identity(row)
        aliases = [key, row.catalog_title]
        title_alias = notice_title_alias(row.catalog_title)
        if title_alias is not None:
            aliases.append(title_alias)
        if identity is not None:
            aliases.extend(identity.aliases)
        target = NoticeImportTarget(
            id=str(uuid4()),
            document_set_id=document_set_id,
            identity_key=key,
            instrument_type=kind,
            instrument_number=number,
            instrument_year=year,
            instrument_date=issued,
            catalog_title=row.catalog_title,
            aliases=list(dict.fromkeys(aliases)),
            legistation_metadata={
                "validity_status": "Mülga",
                "repeal_note": row.catalog_title,
                "validity_source": workbook_path.name,
            },
            provenance={
                "source_filename": workbook_path.name,
                "source_sha256": workbook_sha256,
                "sheet": sheet,
                "excel_rows": [item.excel_row for item in group],
                "raw_records": [item.model_dump(mode="json") for item in group],
                "identity_basis": "opaque_raw_catalogue_observation",
                "structured_identity_basis": basis,
                "prior_review_reason": prior_decision.reason,
                "existing_source_candidates": prior_decision.existing_source_ids,
                "annotation_only": True,
                "links_existing_source": False,
                "repeal_date_supplied": False,
            },
        )
        already = registry_by_key.get(key)
        if already is not None:
            if already.model_dump(exclude={"id"}, mode="json") != target.model_dump(
                exclude={"id"}, mode="json"
            ):
                raise ValueError(
                    "Catalogue observation already exists with different fields"
                )
            classification, reason = (
                "already_registered",
                "catalogue_observation_already_imported_unchanged",
            )
        else:
            targets.append(target)
            classification, reason = (
                "new",
                "independent_name_only_catalogue_observation",
            )
        decisions.append(
            NoticeRowDecision(
                excel_rows=[item.excel_row for item in group],
                identity_key=key,
                classification=classification,
                reason=reason,
                catalog_title=row.catalog_title,
                existing_source_ids=prior_decision.existing_source_ids,
                raw_records=group,
            )
        )
    return SourceNoticeManifest(
        workbook_sha256=workbook_sha256,
        source_path=str(workbook_path),
        sheet=sheet,
        document_set_id=document_set_id,
        inventory_fingerprint=inventory_fingerprint(sources),
        inventory_count=len(sources),
        identity_heads_fingerprint=canonical_digest(
            [head.model_dump(mode="json") for head in identity_heads]
        ),
        identity_heads_count=len(identity_heads),
        raw_row_count=len(rows),
        targets=targets,
        decisions=sorted(decisions, key=lambda item: min(item.excel_rows)),
        admission_policy="catalogue_observations",
        prior_manifest_sha256=prior_manifest_sha256,
        registry_fingerprint=notice_registry_fingerprint(registry),
        baseline_notices=registry,
    )
