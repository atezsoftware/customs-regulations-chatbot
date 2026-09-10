from hashlib import sha256
from io import BytesIO
from unittest.mock import MagicMock
from uuid import uuid4

import pytest

from onyx.llm.interfaces import LLMConfig
from onyx.regulatory.amendments.annexes import corrections
from onyx.regulatory.amendments.annexes.extraction import extract_annex_structure
from onyx.regulatory.amendments.annexes.models import (
    AnnexChangeDraft,
    AnnexCorrectionReconciliation,
    AnnexElementCorrection,
    AnnexOriginalEvidence,
    AnnexReviewEvidence,
)


def office_original(format: str) -> tuple[bytes, str]:
    stream = BytesIO()
    if format == "docx":
        from docx import Document

        document = Document()
        document.add_paragraph("EK-1")
        document.add_paragraph("Printed source value")
        document.add_paragraph("Independent original context")
        document.save(stream)
        mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    else:
        from openpyxl import Workbook
        from openpyxl.worksheet.worksheet import Worksheet

        workbook = Workbook()
        sheet = workbook.active
        assert isinstance(sheet, Worksheet)
        sheet.title = "EK-1"
        sheet["A1"] = "Printed source value"
        sheet["A2"] = "Independent original context"
        workbook.save(stream)
        workbook.close()
        mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return stream.getvalue(), mime


@pytest.mark.parametrize("format", ["docx", "xlsx"])
@pytest.mark.parametrize("selected_view", [False, True])
@pytest.mark.parametrize(
    "source_problem",
    [
        None,
        "missing_original",
        "wrong_locator",
        "ambiguous_original",
        "unreadable_original",
    ],
)
def test_office_correction_supplies_independent_bound_source_or_fails_closed(
    monkeypatch: pytest.MonkeyPatch,
    format: str,
    source_problem: str | None,
    selected_view: bool,
) -> None:
    content, mime = office_original(format)
    original = extract_annex_structure(content, mime)
    position = next(
        index
        for index, element in enumerate(original.elements)
        if element.text == "Printed source value"
    )
    raw = original.model_copy(deep=True)
    for element in raw.elements:
        if element.text != "EK-1":
            element.text = "OCR transcription error"
    if source_problem == "wrong_locator":
        raw.elements[position].locator.path = "unverified/position"
    evidence = AnnexReviewEvidence(
        id=uuid4(),
        side="new",
        kind="original",
        file_id="frozen-office",
        sha256=sha256(content).hexdigest(),
        mime_type=mime,
        byte_count=len(content),
        parent_file_id="original-office",
        parent_sha256=sha256(content).hexdigest(),
    )
    if selected_view:
        from onyx.regulatory.amendments.annexes.evidence import (
            select_annex_evidence_view,
        )

        raw = select_annex_evidence_view(
            extraction=raw,
            original=AnnexOriginalEvidence(
                file_id="original-office",
                sha256=evidence.sha256,
                mime_type=mime,
                available=True,
            ),
            annex_label="EK-1",
            canonical_labels=["EK-1"],
            canonical_chunk_ids=["legal-scope"],
        )
    if source_problem == "unreadable_original":
        original.elements[position].status = "unreadable"
        monkeypatch.setattr(
            "onyx.regulatory.amendments.annexes.extraction.extract_annex_structure",
            lambda _content, _mime: original,
        )
    draft = AnnexChangeDraft(
        instruction_indices=[0],
        instruction_texts=["EK-1"],
        annex_label="EK-1",
        raw_new_extraction=raw,
        new_extraction=raw,
        evidence=[] if source_problem == "missing_original" else [evidence],
    )
    if source_problem == "ambiguous_original":
        draft.evidence.append(
            evidence.model_copy(update={"id": uuid4(), "file_id": "other-original"})
        )
    edit = AnnexElementCorrection(
        position=position,
        before_text="OCR transcription error",
        corrected_text="Printed source value",
        reason="Native source correction",
    )
    store = MagicMock()
    store.read_file.side_effect = lambda _id: BytesIO(content)
    monkeypatch.setattr(corrections, "get_default_file_store", lambda: store)
    transport = MagicMock(
        return_value=AnnexCorrectionReconciliation(
            supported=True, rationale="Native source agrees"
        )
    )
    monkeypatch.setattr(corrections, "generate_structured", transport)
    llm = MagicMock()
    llm.config = LLMConfig(
        model_provider="configured",
        model_name="review",
        temperature=0,
        max_input_tokens=10000,
    )
    if source_problem:
        with pytest.raises(ValueError, match="correction source"):
            corrections.reconcile_corrections(draft=draft, corrections=[edit], llm=llm)
        transport.assert_not_called()
    else:
        receipt = corrections.reconcile_corrections(
            draft=draft, corrections=[edit], llm=llm
        )
        prompt = transport.call_args.kwargs["user_prompt"]
        native_input = prompt.split("Original native source:\n", 1)[1].split(
            "\nProposed corrections:", 1
        )[0]
        assert prompt.startswith(
            "Derived transcription under correction (not original source authority):\n"
        )
        system = transport.call_args.kwargs["system_prompt"]
        assert (
            "A mismatch between the derived transcription and an original is the error being corrected"
            in system
        )
        assert "If the original sources themselves conflict" in system
        assert "Printed source value" in native_input
        assert "OCR transcription error" not in native_input
        assert evidence.sha256 in native_input
        assert original.elements[position].locator.model_dump_json() in native_input
        assert receipt.supported and receipt.input_sha256
